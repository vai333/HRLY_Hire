"""CV text extraction and the PII shield.

Everything that identifies a person — name, email, phone, links, street address, college names and
protected attributes (DOB, gender, marital status, religion, caste, nationality) — is removed here,
before any text leaves the backend for the LLM. The extracted contact details are returned separately
so they can be stored in the database and used for the email, but they never reach Gemini.
"""
from __future__ import annotations

import io
import logging
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional

SUPPORTED_EXTENSIONS = {"pdf", "docx", "txt"}
logging.getLogger("pdfminer").setLevel(logging.ERROR)


# ---------------------------------------------------------------- extraction

def _char_key(c: dict) -> tuple:
    return (c["text"], round(c["x0"], 2), round(c["top"], 2), c.get("fontname"), round(c.get("size", 0), 2))


def _drop_overlays(page):
    """Some resumes stack two renderings of the same line (e.g. "ROHAN MEHTA" over "Rohan Mehta"), which
    extracts as interleaved letters. Keep the first-drawn glyph and drop later ones that overlap it."""
    lines: Dict[int, List[dict]] = {}
    drop = set()
    for c in page.chars:
        if not c["text"].strip():
            continue
        row = lines.setdefault(round(c["top"]), [])
        w = max(c["x1"] - c["x0"], 0.1)
        for k in row:
            overlap = min(c["x1"], k["x1"]) - max(c["x0"], k["x0"])
            if overlap > 0.5 * min(w, k["x1"] - k["x0"]) and abs(c["top"] - k["top"]) < 1:
                drop.add(_char_key(c))
                break
        else:
            row.append(c)
    if not drop:
        return page
    return page.filter(lambda o: o.get("object_type") != "char" or _char_key(o) not in drop)


def extract_text(filename: str, data: bytes) -> str:
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pdf":
        import pdfplumber

        with pdfplumber.open(io.BytesIO(data)) as pdf:
            # dedupe_chars drops "fake bold" overprinting — the same glyph drawn twice, sometimes in two fonts
            # (e.g. "RROohHaAnN"), so ignore font/size when matching duplicates.
            text = "\n".join(
                (_drop_overlays(page.dedupe_chars(tolerance=1.5, extra_attrs=())).extract_text(x_tolerance=1.5) or "")
                for page in pdf.pages
            )
    elif ext == "docx":
        import docx

        doc = docx.Document(io.BytesIO(data))
        chunks: List[str] = []
        for section in doc.sections:  # names often live in the page header
            chunks += [p.text for p in section.header.paragraphs]
        chunks += [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                chunks.append(" | ".join(dict.fromkeys(c.text.strip() for c in row.cells if c.text.strip())))
        text = "\n".join(chunks)
    elif ext == "txt":
        text = data.decode("utf-8", errors="replace")
    else:
        raise ValueError("Unsupported file type. Upload a PDF or DOCX.")

    text = text.replace(" ", " ").replace("\r", "")
    text = re.sub(r"[ \t]+", " ", text)
    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ---------------------------------------------------------------- redaction

EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
URL_RE = re.compile(r"(?:https?://|www\.)\S+|\b(?:linkedin\.com|github\.com|behance\.net|medium\.com)/\S*", re.I)
PHONE_RE = re.compile(r"(?<![\w])\+?\(?\d[\d\s().-]{8,}\d(?![\w])")
PIN_RE = re.compile(r"\b[1-9]\d{2}\s?\d{3}\b")
SENSITIVE_LINE_RE = re.compile(
    r"^[ \t•\-*]*(?:date of birth|d\.?o\.?b\.?|age|gender|sex|marital status|nationality|religion|caste|"
    r"father'?s name|mother'?s name|spouse|passport(?: no\.?| number)?|aadhaa?r|pan(?: no\.?| number)?|"
    r"(?:permanent |current |residential |home )?address)\s*[:\-–|].*$",
    re.I | re.M,
)
HONORIFIC_RE = re.compile(r"\b(?:Mr|Mrs|Ms|Miss|Mx|Shri|Smt|Kumari)\.?\s+(?=[A-Z])")
NAME_LABEL_RE = re.compile(r"^\s*(?:full\s+)?name\s*[:\-–]\s*(.+)$", re.I | re.M)

_INST_WORD = r"[A-Z][\w&.'’-]*"
_INST_KEYWORD = (
    r"(?:University|College|Institute|Polytechnic|Academy|Vidyalaya|Vidyapeeth|Vidyapith|Mahavidyalaya|"
    r"Business School|Law School|Medical School|School of " + _INST_WORD + r"(?:\s+(?:and\s+|&\s+)?" + _INST_WORD + r"){0,3})"
)
INSTITUTION_RE = re.compile(
    r"(?:\b" + _INST_WORD + r"\s+(?:(?:of|and|for|&)\s+)?){0,6}" + _INST_KEYWORD
    + r"(?:\s+(?:of|for)\s+" + _INST_WORD + r"(?:\s+(?:and\s+|&\s+)?" + _INST_WORD + r"){0,4})?"
)
INSTITUTION_ACRONYM_RE = re.compile(
    r"\b(?:IIT|IIM|NIT|IIIT|BITS|ISB|XLRI|SPJIMR|NMIMS|JBIMS|IISc|VJTI|DTU|NSIT|SRCC|LSR|SIBM|IIFT|IMT|MICA|TISS|"
    r"XIMB|XIM|IRMA|TAPMI|GLIM|SCMHRD|KJSIM|IMI|NITIE|LBSIM|BIMTECH|MDI|IIMs|IITs|NITs)\b"
    r"(?:[\s,–-]+(?!Graduate|Alumni|Alumnus)[A-Z][a-z]+)?"
)

CITIES = [
    "Navi Mumbai", "Mumbai", "Bombay", "Thane", "Pune", "Bengaluru", "Bangalore", "New Delhi", "Delhi", "Gurugram",
    "Gurgaon", "Noida", "Hyderabad", "Chennai", "Kolkata", "Ahmedabad", "Jaipur", "Kochi", "Cochin", "Chandigarh",
    "Indore", "Nagpur", "Surat", "Vadodara", "Lucknow", "Coimbatore", "Visakhapatnam", "Goa", "Nashik", "Bhopal",
    "Dubai", "Singapore", "London", "Rotterdam", "Hamburg", "New York", "San Francisco", "Toronto", "Sydney",
]
_NAME_STOPWORDS = {
    "resume", "curriculum", "vitae", "cv", "profile", "summary", "product", "manager", "senior", "engineer",
    "experience", "contact", "objective", "education", "skills", "page", "india", "mumbai", "consultant",
    "experienced", "lead", "head", "director", "analyst", "associate", "specialist", "executive", "operations",
    "logistics", "freight", "supply", "chain", "developer", "designer", "marketing", "sales", "business",
    "founder", "intern", "professional", "about", "me", "personal", "details", "career", "work", "history",
    "projects", "achievements", "technical", "software", "data", "scientist", "management", "owner", "results",
}


@dataclass
class RedactionResult:
    text: str
    name: Optional[str]
    email: Optional[str]
    phone: Optional[str]
    city: Optional[str]
    report: Dict[str, int] = field(default_factory=dict)


def _guess_name(lines: List[str]) -> Optional[str]:
    for line in lines:
        m = NAME_LABEL_RE.match(line)
        if m:
            return m.group(1).split("|")[0].strip() or None
    for line in lines[:6]:
        seg = re.split(r"\s*[|•·,]\s*|\s{3,}| – | - ", line.strip())[0].strip()
        seg = HONORIFIC_RE.sub("", seg)
        tokens = seg.split()
        if not 2 <= len(tokens) <= 4:
            continue
        if not all(re.fullmatch(r"[A-Za-z][A-Za-z.'-]*", t) for t in tokens):
            continue
        if any(t.lower().strip(".") in _NAME_STOPWORDS for t in tokens):
            continue
        # Acronyms like PM / SPM / CTO in a mixed-case line ("Experienced PM"). All-caps names ("ROHAN RAO")
        # and initials ("R. K. Sharma") are still allowed.
        if not all(t.isupper() for t in tokens) and any(t.isupper() and 2 <= len(t.strip(".")) <= 3 for t in tokens):
            continue
        if all(t[0].isupper() for t in tokens):
            return " ".join(t.title() if t.isupper() else t for t in tokens)
    return None


def _sub_count(pattern: re.Pattern, repl, text: str):
    return pattern.subn(repl, text)


def detect_contact(text: str) -> Dict[str, Optional[str]]:
    """Best-effort name / email / phone from the CV, for pre-filling the applicant form."""
    lines = [l for l in text.splitlines() if l.strip()]
    emails = EMAIL_RE.findall(text)
    phone = None
    for m in PHONE_RE.finditer(text):
        if len(re.sub(r"\D", "", m.group())) >= 10:
            phone = m.group().strip()
            break
    return {"name": _guess_name(lines), "email": emails[0] if emails else None, "phone": phone}


def _phone_pattern(phone: str) -> Optional[re.Pattern]:
    digits = re.sub(r"\D", "", phone)[-10:]  # match on the last 10 digits, whatever the formatting
    if len(digits) < 8:
        return None
    return re.compile(r"(?:\+?\d{1,3}[\s().-]*)?" + r"[\s().-]*".join(digits))


def redact(text: str, name: Optional[str] = None, email: Optional[str] = None, phone: Optional[str] = None) -> RedactionResult:
    """Strip PII. `name` / `email` / `phone`, when given (e.g. confirmed by the applicant), take precedence
    over detection and are removed wherever they appear."""
    report: Dict[str, int] = {}
    lines = [l for l in text.splitlines() if l.strip()]
    header = "\n".join(lines[:12])

    detected = detect_contact(text)
    name = (name or "").strip() or detected["name"]
    email = (email or "").strip() or detected["email"]
    phone = (phone or "").strip() or detected["phone"]

    city = None
    city_scope = header + "\n" + "\n".join(re.findall(r"(?im)^\s*(?:location|based in|city)\s*[:\-–].*$", text))
    for c in CITIES:
        if re.search(rf"\b{re.escape(c)}\b", city_scope, re.I):
            city = "Mumbai" if c == "Bombay" else c
            break

    out = text
    out, n = _sub_count(SENSITIVE_LINE_RE, "[PERSONAL DETAIL REMOVED]", out)
    report["personal_details"] = n
    out, n = _sub_count(EMAIL_RE, "[EMAIL]", out)
    report["emails"] = n
    if email:
        out, n = re.subn(re.escape(email), "[EMAIL]", out, flags=re.I)
        report["emails"] += n
    out, n = _sub_count(URL_RE, "[LINK]", out)
    report["links"] = n

    phones = 0

    def _phone(m: re.Match) -> str:
        nonlocal phones
        if len(re.sub(r"\D", "", m.group())) >= 10:
            phones += 1
            return "[PHONE]"
        return m.group()

    out = PHONE_RE.sub(_phone, out)
    pat = _phone_pattern(phone) if phone else None
    if pat:
        out, n = pat.subn("[PHONE]", out)
        phones += n
    report["phones"] = phones

    # Street addresses: header lines that carry a PIN code.
    out_lines = out.splitlines()
    addr = 0
    for i, line in enumerate(out_lines[:15]):
        if PIN_RE.search(line) and ("," in line or any(c.lower() in line.lower() for c in CITIES)):
            out_lines[i] = "[ADDRESS]"
            addr += 1
    out = "\n".join(out_lines)
    report["addresses"] = addr

    out, n_inst = _sub_count(INSTITUTION_RE, "[INSTITUTION]", out)
    out, n_acr = _sub_count(INSTITUTION_ACRONYM_RE, "[INSTITUTION]", out)
    report["institutions"] = n_inst + n_acr

    out = HONORIFIC_RE.sub("", out)

    n_name = 0
    # Strip both the confirmed name and the one printed on the resume (they can differ, e.g. a typo in the form).
    names = [n for n in dict.fromkeys([name, detected["name"]]) if n]
    for full in names:
        out, n = re.subn(re.escape(full), "[CANDIDATE]", out, flags=re.I)
        n_name += n
    tokens = {t for full in names for t in re.split(r"[\s.]+", full) if len(t) >= 3}
    for token in sorted(tokens, key=len, reverse=True):
        out, n = re.subn(rf"\b{re.escape(token)}\b", "[CANDIDATE]", out, flags=re.I)
        n_name += n
    out = re.sub(r"\[CANDIDATE\](?:\s+\[CANDIDATE\])+", "[CANDIDATE]", out)
    report["name_mentions"] = n_name

    loc_line = f"Candidate location (city only, system-extracted): {city or 'not stated in resume header'}"
    out = loc_line + "\n\n" + out.strip()
    return RedactionResult(out, name, email, phone, city, report)


def contact_mismatch_note(name: Optional[str], email: Optional[str]) -> Optional[str]:
    """Rubric A5 asks to flag contact details that don't match the name. The AI never sees either, so we check here."""
    if not name or not email:
        return None
    handle = re.sub(r"[^a-z]", "", email.split("@")[0].lower())
    tokens = [t.lower() for t in re.split(r"[\s.'-]+", name) if len(t) >= 3]
    if tokens and not any(t in handle for t in tokens) and not any(handle.startswith(t[0]) for t in tokens):
        return "Contact check: email handle does not resemble the candidate name — confirm at screen."
    return None
